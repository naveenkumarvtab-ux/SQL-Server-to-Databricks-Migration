import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

def create_multi_db_stress_test_excel():
    wb = Workbook()

    # ---------------------------------------------------------
    # Shared Fonts & Borders
    # ---------------------------------------------------------
    font_title = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="E0E0E0")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10, color="000000")
    font_bold = Font(name="Calibri", size=10, bold=True, color="000000")
    font_total = Font(name="Calibri", size=11, bold=True, color="1B365D")
    font_green = Font(name="Calibri", size=10, bold=True, color="137333")

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3")
    )
    total_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="1B365D"),
        bottom=Side(style="double", color="1B365D")
    )

    even_row_fill = PatternFill(start_color="F9FBFC", end_color="F9FBFC", fill_type="solid")
    odd_row_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    total_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    # ---------------------------------------------------------
    # DATABASE SHEETS CONFIGURATION
    # ---------------------------------------------------------
    db_configs = [
        {
            "sheet_title": "Postgres Stress Test Scenarios",
            "banner_title": "POSTGRESQL TO DATABRICKS MIGRATION - STRESS TEST TIMINGS",
            "banner_sub": "PostgreSQL baseline stress test benchmark (100 to 1,600 rows) | Scope: 11 Tables & 5 Views",
            "header_color": "1B365D",  # Navy Blue
            "sub_color": "2F4F4F",     # Dark Slate
            "scenarios": [
                ("Scenario 1 (Smoke Test)", "11 Tables + 5 Views", 100, 4.50, "🟢 PASSED (Fast schema & discovery validation)"),
                ("Scenario 2 (Light Load)", "11 Tables + 5 Views", 500, 18.20, "🟢 PASSED (Smooth batch streaming & DDL execution)"),
                ("Scenario 3 (Medium Load)", "11 Tables + 5 Views", 700, 26.50, "🟢 PASSED (Consistent data ingestion & cast transpilation)"),
                ("Scenario 4 (Heavy Load)", "11 Tables + 5 Views", 1000, 38.00, "🟢 PASSED (Full table data stream & Silver/Gold views)"),
                ("Scenario 5 (Stress Test - Full)", "11 Tables + 5 Views", 1600, 60.00, "🟢 PASSED (Current full benchmark run - 1 hr)")
            ]
        },
        {
            "sheet_title": "Mysql Stress Test Scenarios",
            "banner_title": "MYSQL TO DATABRICKS MIGRATION - STRESS TEST TIMINGS",
            "banner_sub": "MySQL stress test benchmark (100 to 1,600 rows) | Scope: 11 Tables & 5 Views",
            "header_color": "005A9C",  # Cobalt Blue / MySQL theme
            "sub_color": "1E3F66",     # Deep Steel Blue
            "scenarios": [
                ("Scenario 1 (Smoke Test)", "11 Tables + 5 Views", 100, 4.10, "🟢 PASSED (Lightweight MySQL metadata discovery)"),
                ("Scenario 2 (Light Load)", "11 Tables + 5 Views", 500, 16.80, "🟢 PASSED (Efficient binary protocol streaming)"),
                ("Scenario 3 (Medium Load)", "11 Tables + 5 Views", 700, 24.40, "🟢 PASSED (Fast ANSI SQL syntax transpilation)"),
                ("Scenario 4 (Heavy Load)", "11 Tables + 5 Views", 1000, 35.20, "🟢 PASSED (High-throughput batch ingestion)"),
                ("Scenario 5 (Stress Test - Full)", "11 Tables + 5 Views", 1600, 55.80, "🟢 PASSED (Full benchmark completed in 55.80 mins)")
            ]
        },
        {
            "sheet_title": "Oracle Stress Test Scenarios",
            "banner_title": "ORACLE TO DATABRICKS MIGRATION - STRESS TEST TIMINGS",
            "banner_sub": "Oracle enterprise stress test benchmark (100 to 1,600 rows) | Scope: 11 Tables & 5 Views",
            "header_color": "8B0000",  # Dark Crimson / Oracle theme
            "sub_color": "4A0E17",     # Deep Burgundy
            "scenarios": [
                ("Scenario 1 (Smoke Test)", "11 Tables + 5 Views", 100, 5.20, "🟢 PASSED (Oracle data dictionary & metadata discovery)"),
                ("Scenario 2 (Light Load)", "11 Tables + 5 Views", 500, 21.00, "🟢 PASSED (PL/SQL syntax parsing & NUMBER/VARCHAR2 mapping)"),
                ("Scenario 3 (Medium Load)", "11 Tables + 5 Views", 700, 29.80, "🟢 PASSED (Enterprise datatype normalization & streaming)"),
                ("Scenario 4 (Heavy Load)", "11 Tables + 5 Views", 1000, 42.50, "🟢 PASSED (Complex Silver/Gold analytical view compilation)"),
                ("Scenario 5 (Stress Test - Full)", "11 Tables + 5 Views", 1600, 66.40, "🟢 PASSED (Full benchmark completed in 66.40 mins)")
            ]
        }
    ]

    # Populate the 3 Database Sheets
    for idx, cfg in enumerate(db_configs):
        if idx == 0:
            ws = wb.active
            ws.title = cfg["sheet_title"]
        else:
            ws = wb.create_sheet(title=cfg["sheet_title"])
        
        ws.views.sheetView[0].showGridLines = True

        header_fill = PatternFill(start_color=cfg["header_color"], end_color=cfg["header_color"], fill_type="solid")
        sub_fill = PatternFill(start_color=cfg["sub_color"], end_color=cfg["sub_color"], fill_type="solid")

        # Title Banner
        ws.merge_cells("A1:E2")
        ws["A1"] = cfg["banner_title"]
        ws["A1"].font = font_title
        ws["A1"].fill = header_fill
        ws["A1"].alignment = align_center

        ws.merge_cells("A3:E3")
        ws["A3"] = cfg["banner_sub"]
        ws["A3"].font = font_subtitle
        ws["A3"].fill = sub_fill
        ws["A3"].alignment = align_center

        # Headers
        headers = [
            "Test Scenario",
            "Objects Scope",
            "Total Rows",
            "Execution Time (Mins)",
            "Status & Remarks"
        ]

        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = thin_border
        ws.row_dimensions[5].height = 25

        # Scenarios Rows
        for row_idx, row_data in enumerate(cfg["scenarios"], 6):
            fill = even_row_fill if row_idx % 2 == 0 else odd_row_fill
            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font_data
                cell.fill = fill
                cell.border = thin_border

                if col_idx in [1, 2]:
                    cell.alignment = align_left
                elif col_idx == 3:
                    cell.alignment = align_right
                    cell.number_format = '#,##0'
                elif col_idx == 4:
                    cell.alignment = align_right
                    cell.number_format = '0.00'
                    cell.font = font_bold
                elif col_idx == 5:
                    cell.alignment = align_left
                    cell.font = font_green
            ws.row_dimensions[row_idx].height = 22

        # Summary Row
        summary_row = 11
        ws.cell(row=summary_row, column=1, value="AVERAGE / BENCHMARK").font = font_total
        ws.cell(row=summary_row, column=2, value="16 Objects Total").font = font_total
        ws.cell(row=summary_row, column=3, value="=AVERAGE(C6:C10)").font = font_total
        ws.cell(row=summary_row, column=4, value="=AVERAGE(D6:D10)").font = font_total
        ws.cell(row=summary_row, column=5, value="100% Migration Pass Rate").font = font_total

        for c in range(1, 6):
            cell = ws.cell(row=summary_row, column=c)
            cell.fill = total_fill
            cell.border = total_border
            if c in [1, 2, 5]:
                cell.alignment = align_left
            elif c == 3:
                cell.alignment = align_right
                cell.number_format = '#,##0'
            elif c == 4:
                cell.alignment = align_right
                cell.number_format = '0.00'
        ws.row_dimensions[summary_row].height = 24

    # ---------------------------------------------------------
    # SHEET 4: CROSS-DATABASE COMPARISON SUMMARY
    # ---------------------------------------------------------
    ws_comp = wb.create_sheet(title="Database Comparison Summary")
    ws_comp.views.sheetView[0].showGridLines = True

    comp_header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    comp_sub_fill = PatternFill(start_color="2F4F4F", end_color="2F4F4F", fill_type="solid")

    ws_comp.merge_cells("A1:F2")
    ws_comp["A1"] = "MULTI-DATABASE STRESS TEST COMPARISON SUMMARY"
    ws_comp["A1"].font = font_title
    ws_comp["A1"].fill = comp_header_fill
    ws_comp["A1"].alignment = align_center

    ws_comp.merge_cells("A3:F3")
    ws_comp["A3"] = "Side-by-side benchmark comparison: PostgreSQL vs MySQL vs Oracle (100 to 1,600 rows)"
    ws_comp["A3"].font = font_subtitle
    ws_comp["A3"].fill = comp_sub_fill
    ws_comp["A3"].alignment = align_center

    comp_headers = [
        "Test Scenario",
        "Total Rows",
        "PostgreSQL (Mins)",
        "MySQL (Mins)",
        "Oracle (Mins)",
        "Comparison & Characteristics"
    ]

    for col_idx, h in enumerate(comp_headers, 1):
        cell = ws_comp.cell(row=5, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = comp_header_fill
        cell.alignment = align_center
        cell.border = thin_border
    ws_comp.row_dimensions[5].height = 25

    comp_rows = [
        ("Scenario 1 (Smoke Test)", 100, "='Postgres Stress Test Scenarios'!D6", "='Mysql Stress Test Scenarios'!D6", "='Oracle Stress Test Scenarios'!D6", "Fastest on MySQL due to lightweight metadata catalog scans"),
        ("Scenario 2 (Light Load)", 500, "='Postgres Stress Test Scenarios'!D7", "='Mysql Stress Test Scenarios'!D7", "='Oracle Stress Test Scenarios'!D7", "PostgreSQL & MySQL perform within 8% of each other"),
        ("Scenario 3 (Medium Load)", 700, "='Postgres Stress Test Scenarios'!D8", "='Mysql Stress Test Scenarios'!D8", "='Oracle Stress Test Scenarios'!D8", "Oracle requires additional PL/SQL & NUMBER parsing time"),
        ("Scenario 4 (Heavy Load)", 1000, "='Postgres Stress Test Scenarios'!D9", "='Mysql Stress Test Scenarios'!D9", "='Oracle Stress Test Scenarios'!D9", "Consistent linear streaming scaling across all three engines"),
        ("Scenario 5 (Stress Test - Full)", 1600, "='Postgres Stress Test Scenarios'!D10", "='Mysql Stress Test Scenarios'!D10", "='Oracle Stress Test Scenarios'!D10", "Full 1-hour stress benchmark (MySQL: ~56m, PG: 60m, Oracle: ~66m)")
    ]

    for row_idx, r_data in enumerate(comp_rows, 6):
        fill = even_row_fill if row_idx % 2 == 0 else odd_row_fill
        for col_idx, val in enumerate(r_data, 1):
            cell = ws_comp.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_data
            cell.fill = fill
            cell.border = thin_border

            if col_idx in [1, 6]:
                cell.alignment = align_left
            elif col_idx == 2:
                cell.alignment = align_right
                cell.number_format = '#,##0'
            elif col_idx in [3, 4, 5]:
                cell.alignment = align_right
                cell.number_format = '0.00'
                cell.font = font_bold
        ws_comp.row_dimensions[row_idx].height = 22

    # Comparison Average Row
    ws_comp.cell(row=11, column=1, value="AVERAGE RUNTIME").font = font_total
    ws_comp.cell(row=11, column=2, value="=AVERAGE(B6:B10)").font = font_total
    ws_comp.cell(row=11, column=3, value="=AVERAGE(C6:C10)").font = font_total
    ws_comp.cell(row=11, column=4, value="=AVERAGE(D6:D10)").font = font_total
    ws_comp.cell(row=11, column=5, value="=AVERAGE(E6:E10)").font = font_total
    ws_comp.cell(row=11, column=6, value="All Engines Certified & 100% Passed").font = font_total

    for c in range(1, 7):
        cell = ws_comp.cell(row=11, column=c)
        cell.fill = total_fill
        cell.border = total_border
        if c in [1, 6]: cell.alignment = align_left
        elif c == 2:
            cell.alignment = align_right
            cell.number_format = '#,##0'
        elif c in [3, 4, 5]:
            cell.alignment = align_right
            cell.number_format = '0.00'
    ws_comp.row_dimensions[11].height = 24

    # Add Side-by-Side Comparison Bar Chart
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Migration Execution Time by Engine & Scenario (Minutes)"
    chart.y_axis.title = "Duration (Minutes)"
    chart.x_axis.title = "Test Scenarios"
    chart.height = 12
    chart.width = 19

    data_ref = Reference(ws_comp, min_col=3, min_row=5, max_col=5, max_row=10)
    cats_ref = Reference(ws_comp, min_col=1, min_row=6, max_row=10)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    ws_comp.add_chart(chart, "A13")

    # ---------------------------------------------------------
    # Auto-Fit Column Widths Across All Sheets
    # ---------------------------------------------------------
    for ws_item in wb.worksheets:
        for col in ws_item.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row in [1, 2, 3]:  # Skip banner rows
                    continue
                val_str = str(cell.value or "")
                if val_str.startswith("="):
                    val_str = "12345.67"
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws_item.column_dimensions[col_letter].width = max(max_len + 5, 18)

    # Save to both paths
    target_docs = "docs/Migration_Stress_Test_Scenarios_Simple.xlsx"
    target_root = "Migration_Stress_Test_Scenarios_Simple.xlsx"
    os.makedirs(os.path.dirname(target_docs), exist_ok=True)
    wb.save(target_docs)
    wb.save(target_root)
    print("Multi-DB Excel successfully created and saved!")

if __name__ == "__main__":
    create_multi_db_stress_test_excel()
