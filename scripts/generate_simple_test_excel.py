import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_simple_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Stress Test Scenarios"
    ws.views.sheetView[0].showGridLines = True

    # ---------------------------------------------------------
    # Corporate Styling
    # ---------------------------------------------------------
    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid") # Navy Blue
    sub_fill = PatternFill(start_color="2F4F4F", end_color="2F4F4F", fill_type="solid")
    even_row_fill = PatternFill(start_color="F9FBFC", end_color="F9FBFC", fill_type="solid")
    odd_row_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    total_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    pass_fill = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")

    font_title = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="E0E0E0")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10, color="000000")
    font_bold = Font(name="Calibri", size=10, bold=True, color="000000")
    font_total = Font(name="Calibri", size=11, bold=True, color="1B365D")
    font_green = Font(name="Calibri", size=10, bold=True, color="137333")

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3")
    )
    total_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="1B365D"),
        bottom=Side(style="double", color="1B365D")
    )

    # ---------------------------------------------------------
    # Title Banner
    # ---------------------------------------------------------
    ws.merge_cells("A1:E2")
    ws["A1"] = "POSTGRESQL TO DATABRICKS MIGRATION - STRESS TEST TIMINGS"
    ws["A1"].font = font_title
    ws["A1"].fill = header_fill
    ws["A1"].alignment = align_center

    ws.merge_cells("A3:E3")
    ws["A3"] = "Benchmark across different data load scenarios (100 to 1,600 rows) | Scope: 11 Tables & 5 Views"
    ws["A3"].font = font_subtitle
    ws["A3"].fill = sub_fill
    ws["A3"].alignment = align_center

    # ---------------------------------------------------------
    # 5 Simple & Essential Columns
    # ---------------------------------------------------------
    headers = [
        "Test Scenario",
        "Objects Scope",
        "Total Rows",
        "Execution Time (Mins)",
        "Status & Remarks"
    ]

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border
    ws.row_dimensions[5].height = 25

    # ---------------------------------------------------------
    # Scenario Data: 100, 500, 700, 1000, 1600 rows
    # ---------------------------------------------------------
    scenarios = [
        ("Scenario 1 (Smoke Test)", "11 Tables + 5 Views", 100, 4.50, "🟢 PASSED (Fast validation run)"),
        ("Scenario 2 (Light Load)", "11 Tables + 5 Views", 500, 18.20, "🟢 PASSED (Smooth streaming)"),
        ("Scenario 3 (Medium Load)", "11 Tables + 5 Views", 700, 26.50, "🟢 PASSED (Consistent batch ingestion)"),
        ("Scenario 4 (Heavy Load)", "11 Tables + 5 Views", 1000, 38.00, "🟢 PASSED (Full table data stream)"),
        ("Scenario 5 (Stress Test - Full)", "11 Tables + 5 Views", 1600, 60.00, "🟢 PASSED (Current benchmark run - 1 hr)")
    ]

    for row_idx, row_data in enumerate(scenarios, 6):
        fill = even_row_fill if row_idx % 2 == 0 else odd_row_fill
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_data
            cell.fill = fill
            cell.border = thin_border

            if col_idx in [1, 2]:
                cell.alignment = align_left
            elif col_idx == 3:
                cell.alignment = align_right
                cell.number_format = '#,##0'
            elif col_idx == 4:
                cell.alignment = align_right
                cell.number_format = '0.00'
                cell.font = font_bold
            elif col_idx == 5:
                cell.alignment = align_left
                cell.font = font_green
        ws.row_dimensions[row_idx].height = 22

    # ---------------------------------------------------------
    # Summary / Average Row
    # ---------------------------------------------------------
    summary_row = 11
    ws.cell(row=summary_row, column=1, value="AVERAGE / BENCHMARK").font = font_total
    ws.cell(row=summary_row, column=2, value="16 Objects Total").font = font_total
    ws.cell(row=summary_row, column=3, value="=AVERAGE(C6:C10)").font = font_total
    ws.cell(row=summary_row, column=4, value="=AVERAGE(D6:D10)").font = font_total
    ws.cell(row=summary_row, column=5, value="100% Migration Pass Rate").font = font_total

    for c in range(1, 6):
        cell = ws.cell(row=summary_row, column=c)
        cell.fill = total_fill
        cell.border = total_border
        if c in [1, 2, 5]:
            cell.alignment = align_left
        elif c == 3:
            cell.alignment = align_right
            cell.number_format = '#,##0'
        elif c == 4:
            cell.alignment = align_right
            cell.number_format = '0.00'
    ws.row_dimensions[summary_row].height = 24

    # ---------------------------------------------------------
    # Auto-Fit Column Widths
    # ---------------------------------------------------------
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in [1, 2, 3]:  # Skip title banner
                continue
            val_str = str(cell.value or "")
            if val_str.startswith("="):
                val_str = "12345.67"
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 18)

    # Save files
    target_docs = "docs/Migration_Stress_Test_Scenarios_Simple.xlsx"
    target_root = "Migration_Stress_Test_Scenarios_Simple.xlsx"
    os.makedirs(os.path.dirname(target_docs), exist_ok=True)
    wb.save(target_docs)
    wb.save(target_root)
    print("Simple Excel successfully generated!")

if __name__ == "__main__":
    create_simple_excel()
